import pytest
import pytest_asyncio
from httpx import AsyncClient, ASGITransport
from sqlalchemy.ext.asyncio import create_async_engine, async_sessionmaker, AsyncSession

import zipfile
from app.main import app
from app.core.config import settings
from app.core.database import Base, get_db
import io
from docx import Document as DocxBuilder
from openpyxl import Workbook
from app.services.scanner import EICAR_SIGNATURE


TEST_DATABASE_URL = "sqlite+aiosqlite:///:memory:"

test_engine = create_async_engine(TEST_DATABASE_URL, echo=False)
TestSessionLocal = async_sessionmaker(
    bind=test_engine,
    class_=AsyncSession,
    expire_on_commit=False,
)


async def override_get_db():
    async with TestSessionLocal() as session:
        yield session


app.dependency_overrides[get_db] = override_get_db

UPLOAD_URL = f"{settings.API_V1_STR}/documents/upload"


@pytest_asyncio.fixture(autouse=True)
async def setup_test_database():
    async with test_engine.begin() as conn:
        await conn.run_sync(Base.metadata.create_all)
    yield
    async with test_engine.begin() as conn:
        await conn.run_sync(Base.metadata.drop_all)


@pytest.mark.asyncio
async def test_valid_pdf_upload():
    async with AsyncClient(
        transport=ASGITransport(app=app), base_url="http://test"
    ) as client:
        dummy_content = b"%PDF-1.4\nTest PDF body content"
        files = {"file": ("document.pdf", dummy_content, "application/pdf")}
        headers = {"X-Tenant-ID": "tenant_123"}

        response = await client.post(UPLOAD_URL, files=files, headers=headers)
        assert response.status_code == 201
        data = response.json()
        assert data["file_format"] == "pdf"
        assert data["mime_type"] == "application/pdf"
        assert data["status"] == "UPLOADED"


@pytest.mark.asyncio
async def test_valid_docx_upload():
    async with AsyncClient(
        transport=ASGITransport(app=app), base_url="http://test"
    ) as client:
        dummy_content = b"PK\x03\x04" + b"\x00" * 50 + b"word/document.xml"
        files = {
            "file": (
                "spec.docx",
                dummy_content,
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        }
        headers = {"X-Tenant-ID": "tenant_123"}

        response = await client.post(UPLOAD_URL, files=files, headers=headers)
        assert response.status_code == 201
        data = response.json()
        assert data["file_format"] == "docx"
        assert (
            data["mime_type"]
            == "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        )


@pytest.mark.asyncio
async def test_valid_pptx_upload():
    async with AsyncClient(
        transport=ASGITransport(app=app), base_url="http://test"
    ) as client:
        dummy_content = b"PK\x03\x04" + b"\x00" * 50 + b"ppt/presentation.xml"
        files = {
            "file": (
                "deck.pptx",
                dummy_content,
                "application/vnd.openxmlformats-officedocument.presentationml.presentation",
            )
        }
        headers = {"X-Tenant-ID": "tenant_123"}

        response = await client.post(UPLOAD_URL, files=files, headers=headers)
        assert response.status_code == 201
        data = response.json()
        assert data["file_format"] == "pptx"
        assert (
            data["mime_type"]
            == "application/vnd.openxmlformats-officedocument.presentationml.presentation"
        )


@pytest.mark.asyncio
async def test_valid_xlsx_upload():
    async with AsyncClient(
        transport=ASGITransport(app=app), base_url="http://test"
    ) as client:
        dummy_content = b"PK\x03\x04" + b"\x00" * 50 + b"xl/workbook.xml"
        files = {
            "file": (
                "metrics.xlsx",
                dummy_content,
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        }
        headers = {"X-Tenant-ID": "tenant_123"}

        response = await client.post(UPLOAD_URL, files=files, headers=headers)
        assert response.status_code == 201
        data = response.json()
        assert data["file_format"] == "xlsx"
        assert (
            data["mime_type"]
            == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )


@pytest.mark.asyncio
async def test_invalid_magic_bytes_rejected():
    async with AsyncClient(
        transport=ASGITransport(app=app), base_url="http://test"
    ) as client:
        corrupted_docx = b"NOT_A_ZIP_CONTENT_PLAIN_TEXT"
        files = {"file": ("fake.docx", corrupted_docx, "application/octet-stream")}
        headers = {"X-Tenant-ID": "tenant_123"}

        response = await client.post(UPLOAD_URL, files=files, headers=headers)
        assert response.status_code == 400
        assert "Invalid file signature" in response.json()["detail"]


@pytest.mark.asyncio
async def test_unsupported_file_extension():
    async with AsyncClient(
        transport=ASGITransport(app=app), base_url="http://test"
    ) as client:
        script_file = b"echo 'malicious'"
        files = {"file": ("test.exe", script_file, "application/octet-stream")}
        headers = {"X-Tenant-ID": "tenant_123"}

        response = await client.post(UPLOAD_URL, files=files, headers=headers)
        assert response.status_code == 400
        assert "Unsupported file format" in response.json()["detail"]

@pytest.mark.asyncio
async def test_docx_metadata_extraction():
    # Build a valid in-memory DOCX with an author and a paragraph
    doc = DocxBuilder()
    doc.core_properties.author = "Test Engineer"
    doc.core_properties.title = "Sample Spec"
    doc.add_paragraph("First paragraph content.")
    buf = io.BytesIO()
    doc.save(buf)
    docx_bytes = buf.getvalue()

    async with AsyncClient(
        transport=ASGITransport(app=app), base_url="http://test"
    ) as client:
        files = {
            "file": (
                "document.docx",
                docx_bytes,
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        }
        headers = {"X-Tenant-ID": "tenant_meta"}

        upload_resp = await client.post(UPLOAD_URL, files=files, headers=headers)
        assert upload_resp.status_code == 201
        doc_id = upload_resp.json()["document_id"]

        status_resp = await client.get(
            f"{settings.API_V1_STR}/documents/{doc_id}", headers=headers
        )
        assert status_resp.status_code == 200
        meta = status_resp.json()["metadata_json"]
        assert "extracted_properties" in meta
        props = meta["extracted_properties"]
        assert props.get("author") == "Test Engineer"
        assert props.get("title") == "Sample Spec"
        assert props.get("paragraph_count") == 1


@pytest.mark.asyncio
async def test_xlsx_metadata_extraction():
    # Build a valid in-memory XLSX with custom sheets and creator
    wb = Workbook()
    wb.properties.creator = "Data Analyst"
    ws = wb.active
    ws.title = "Summary"
    wb.create_sheet(title="RawData")
    buf = io.BytesIO()
    wb.save(buf)
    xlsx_bytes = buf.getvalue()

    async with AsyncClient(
        transport=ASGITransport(app=app), base_url="http://test"
    ) as client:
        files = {
            "file": (
                "sheets.xlsx",
                xlsx_bytes,
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        }
        headers = {"X-Tenant-ID": "tenant_meta"}

        upload_resp = await client.post(UPLOAD_URL, files=files, headers=headers)
        assert upload_resp.status_code == 201
        doc_id = upload_resp.json()["document_id"]

        status_resp = await client.get(
            f"{settings.API_V1_STR}/documents/{doc_id}", headers=headers
        )
        assert status_resp.status_code == 200
        meta = status_resp.json()["metadata_json"]
        props = meta["extracted_properties"]
        assert props.get("author") == "Data Analyst"
        assert props.get("sheet_count") == 2
        assert props.get("sheet_names") == ["Summary", "RawData"]
        

@pytest.mark.asyncio
async def test_malware_eicar_signature_rejected():
    async with AsyncClient(
        transport=ASGITransport(app=app), base_url="http://test"
    ) as client:
        # Construct a PDF header that contains the EICAR test signature
        malicious_pdf = b"%PDF-1.4\n" + EICAR_SIGNATURE
        files = {"file": ("infected.pdf", malicious_pdf, "application/pdf")}
        headers = {"X-Tenant-ID": "tenant_sec"}

        response = await client.post(UPLOAD_URL, files=files, headers=headers)
        assert response.status_code == 422
        assert "Security quarantine" in response.json()["detail"]
        assert "Eicar-Test-Signature" in response.json()["detail"]

@pytest.mark.asyncio
async def test_valid_zip_archive_upload():
    # Build a clean in-memory zip containing two text/doc entries
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zf:
        zf.writestr("manifest.txt", "File list content")
        zf.writestr("notes.txt", "Onboarding notes")
    zip_bytes = buf.getvalue()

    async with AsyncClient(
        transport=ASGITransport(app=app), base_url="http://test"
    ) as client:
        files = {"file": ("bundle.zip", zip_bytes, "application/zip")}
        headers = {"X-Tenant-ID": "tenant_zip"}

        response = await client.post(UPLOAD_URL, files=files, headers=headers)
        assert response.status_code == 201
        data = response.json()
        assert data["file_format"] == "zip"

        # Check status endpoint metadata for extracted manifest
        status_resp = await client.get(
            f"{settings.API_V1_STR}/documents/{data['document_id']}", headers=headers
        )
        assert status_resp.status_code == 200
        meta = status_resp.json()["metadata_json"]
        props = meta["extracted_properties"]
        assert props["total_unpacked_files"] == 2


@pytest.mark.asyncio
async def test_zip_path_traversal_rejected():
    # Build a zip containing a malicious path traversal attempt (Zip Slip)
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as zf:
        zf.writestr("../../etc/passwd", "malicious traversal content")
    traversal_bytes = buf.getvalue()

    async with AsyncClient(
        transport=ASGITransport(app=app), base_url="http://test"
    ) as client:
        files = {"file": ("exploit.zip", traversal_bytes, "application/zip")}
        headers = {"X-Tenant-ID": "tenant_zip"}

        response = await client.post(UPLOAD_URL, files=files, headers=headers)
        assert response.status_code == 400
        assert "Path traversal detected" in response.json()["detail"]


@pytest.mark.asyncio
async def test_zip_containing_malware_quarantined():
    # Build a zip containing an EICAR infected file inside
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as zf:
        zf.writestr("clean.txt", "Safe contents")
        zf.writestr("virus.bin", EICAR_SIGNATURE)
    infected_zip = buf.getvalue()

    async with AsyncClient(
        transport=ASGITransport(app=app), base_url="http://test"
    ) as client:
        files = {"file": ("infected.zip", infected_zip, "application/zip")}
        headers = {"X-Tenant-ID": "tenant_zip"}

        response = await client.post(UPLOAD_URL, files=files, headers=headers)
        assert response.status_code == 422
        assert "Security quarantine" in response.json()["detail"]